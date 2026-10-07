import os
import time
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

from gateways import render_gateways

# ============================================================
# CONFIGURACIÓN
# ============================================================
st.set_page_config(
    page_title="Elite Flower Assistant",
    page_icon="🌸",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# Carpeta del repo donde van los documentos (junto a app.py)
DOCS_DIR = Path(__file__).parent / "documentos"
SUPPORTED_EXT = {".pdf", ".docx", ".xlsx", ".xls", ".csv", ".txt"}

DEFAULT_MODEL = "gemini-flash-latest"
MAX_CHARS_PER_FILE = 120_000
MAX_CHARS_TOTAL = 600_000


def get_secret(name: str, default=None):
    """Lee de st.secrets o de variables de entorno, sin romperse si no existen."""
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass
    return os.environ.get(name, default)


API_KEY = get_secret("GEMINI_API_KEY")
MODEL = get_secret("GEMINI_MODEL", DEFAULT_MODEL)

# Si el modelo principal está saturado (503) o no existe (404), se prueban estos en orden.
FALLBACK_MODELS = ["gemini-3.5-flash", "gemini-3.5-flash-lite", "gemini-2.5-flash"]
MODELS_TO_TRY = [MODEL] + [m for m in FALLBACK_MODELS if m != MODEL]
RETRIES_PER_MODEL = 3

# ============================================================
# ESTILO — DARK / NEON (sin barra lateral)
# ============================================================
st.markdown(
    """
<style>
.stApp {
    background:
        radial-gradient(circle at 78% 12%, rgba(255, 50, 170, .10), transparent 28%),
        radial-gradient(circle at 15% 75%, rgba(0, 255, 210, .08), transparent 30%),
        #05070b;
    color: #eef2ff;
}
.block-container { max-width: 1380px; padding-top: 1.5rem; padding-bottom: 4rem; }
/* Oculta la barra lateral y su botón */
section[data-testid="stSidebar"],
div[data-testid="stSidebarCollapsedControl"],
div[data-testid="collapsedControl"] { display: none !important; }

.hero { padding: 24px 0 18px 0; text-align: center; }
.hero-title {
    font-size: 38px; font-weight: 850; color: #fff;
    text-shadow: 0 0 12px rgba(255, 55, 190, .35); margin: 0;
}
.hero-sub { color: #929bb0; font-size: 15px; margin-top: 6px; }
.neon { color: #ff4fc3; text-shadow: 0 0 12px rgba(255, 79, 195, .45); }
.topbar {
    display: flex; align-items: center; justify-content: space-between;
    gap: 16px; padding: 12px 0 18px;
    border-bottom: 1px solid rgba(255,255,255,.08); margin-bottom: 18px;
}
.topbar-brand { font-size: 13px; font-weight: 800; letter-spacing: .14em; color: #f8fafc; }
.topbar-brand span { color: #ff4fc3; }
.topbar-note { color: #8e9aaa; font-size: 12px; }
.stButton > button {
    background: #0b1018; color: #00ffd5;
    border: 1px solid rgba(0,255,213,.35); border-radius: 10px;
}
.stButton > button:hover {
    border-color: #00ffd5; color: #00ffd5; box-shadow: 0 0 15px rgba(0,255,213,.2);
}
div[data-testid="stChatMessage"] {
    background: rgba(10,14,22,.82);
    border: 1px solid rgba(255,255,255,.07);
    border-radius: 14px;
}
</style>
""",
    unsafe_allow_html=True,
)


# ============================================================
# LECTURA DE DOCUMENTOS (desde la carpeta del repo)
# ============================================================
def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_text(path: Path) -> str:
    """Extrae texto de PDF, Word, Excel, CSV o TXT."""
    ext = path.suffix.lower()
    data = path.read_bytes()
    try:
        if ext == ".pdf":
            reader = PdfReader(BytesIO(data))
            pages = []
            for i, page in enumerate(reader.pages, start=1):
                pages.append(f"[Página {i}]\n{page.extract_text() or ''}")
            text = "\n\n".join(pages)

        elif ext == ".docx":
            doc = Document(BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for t_idx, table in enumerate(doc.tables, start=1):
                parts.append(f"[Tabla {t_idx}]")
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts)

        elif ext in (".xlsx", ".xls"):
            sheets = pd.read_excel(BytesIO(data), sheet_name=None)
            parts = []
            for sheet_name, df in sheets.items():
                parts.append(f"[Hoja: {sheet_name}] ({len(df)} filas)")
                parts.append(df.to_csv(index=False))
            text = "\n".join(parts)

        elif ext == ".csv":
            df = pd.read_csv(BytesIO(data), sep=None, engine="python")
            text = f"({len(df)} filas)\n" + df.to_csv(index=False)

        elif ext == ".txt":
            text = _decode_text(data)

        else:
            return f"[Formato {ext} no soportado]"
    except Exception as e:
        return f"[No se pudo leer el archivo: {e}]"

    if len(text) > MAX_CHARS_PER_FILE:
        text = text[:MAX_CHARS_PER_FILE] + "\n[...contenido truncado por tamaño...]"
    return text


def list_document_files() -> list[Path]:
    """Lista los archivos soportados dentro de documentos/ (incluye subcarpetas)."""
    if not DOCS_DIR.exists():
        return []
    files = []
    for p in sorted(DOCS_DIR.rglob("*")):
        if not p.is_file():
            continue
        if p.name.startswith((".", "~$")):  # ocultos y temporales de Word
            continue
        if p.suffix.lower() in SUPPORTED_EXT:
            files.append(p)
    return files


def files_signature(files: list[Path]) -> tuple:
    """Huella de los archivos: si cambia alguno, se vuelve a leer todo."""
    return tuple((str(p), p.stat().st_mtime, p.stat().st_size) for p in files)


@st.cache_data(show_spinner=False)
def load_documents(signature: tuple) -> str:
    """Devuelve el texto de todos los documentos, listo para dárselo al modelo."""
    blocks, total = [], 0
    for path_str, _, _ in signature:
        path = Path(path_str)
        rel = path.relative_to(DOCS_DIR).as_posix()
        text = extract_text(path)
        if total + len(text) > MAX_CHARS_TOTAL:
            blocks.append(f"=== ARCHIVO: {rel} ===\n[Omitido: se alcanzó el límite total]")
            continue
        blocks.append(f"=== ARCHIVO: {rel} ===\n{text}")
        total += len(text)
    return "\n\n".join(blocks)


# ============================================================
# LLAMADA A GEMINI
# ============================================================
def ask_gemini(history: list[dict], docs_context: str) -> str:
    client = genai.Client(api_key=API_KEY)

    system = (
        "Eres el asistente de Elite Flower, una empresa floricultora. "
        "Responde siempre en español, de forma clara, precisa y concisa. "
    )
    if docs_context:
        system += (
            "Responde basándote en los documentos de la empresa que aparecen abajo. "
            "Si la respuesta no está en ellos, dilo claramente en lugar de inventar. "
            "Cuando sea útil, menciona de qué archivo sacaste la información.\n\n"
            "DOCUMENTOS:\n" + docs_context
        )
    else:
        system += (
            "No tienes documentos de la empresa disponibles. Puedes responder con "
            "conocimiento general, aclarando que no cuentas con información interna."
        )

    # Gemini espera que la conversación empiece con un mensaje del usuario.
    msgs = list(history)
    while msgs and msgs[0]["role"] != "user":
        msgs.pop(0)

    contents = [
        types.Content(
            role="user" if m["role"] == "user" else "model",
            parts=[types.Part(text=m["content"])],
        )
        for m in msgs
    ]

    config = types.GenerateContentConfig(system_instruction=system, temperature=0.3)

    last_error = None
    overload_error = None
    for model_name in MODELS_TO_TRY:
        for attempt in range(RETRIES_PER_MODEL):
            try:
                response = client.models.generate_content(
                    model=model_name, contents=contents, config=config
                )
                return response.text or "No obtuve respuesta del modelo. Intenta reformular la pregunta."
            except Exception as e:
                last_error = e
                msg = str(e)
                if "503" in msg or "UNAVAILABLE" in msg or "429" in msg or "RESOURCE_EXHAUSTED" in msg:
                    overload_error = e
                    if attempt < RETRIES_PER_MODEL - 1:
                        time.sleep(2 * (attempt + 1))
                        continue
                    break  # agotó reintentos: pasa al siguiente modelo
                if "404" in msg or "NOT_FOUND" in msg:
                    break  # el modelo no existe: pasa al siguiente
                raise  # otro error (llave inválida, etc.)

    raise (overload_error or last_error)


def friendly_error(e: Exception) -> str:
    msg = str(e)
    if "503" in msg or "UNAVAILABLE" in msg:
        return "⚠️ El servicio está saturado en este momento. Espera un minuto y vuelve a preguntar."
    if "429" in msg or "RESOURCE_EXHAUSTED" in msg:
        return "⚠️ Se alcanzó el límite de uso. Espera un momento e intenta de nuevo."
    if "404" in msg or "NOT_FOUND" in msg:
        return "⚠️ El modelo de IA configurado no está disponible. Avisa al administrador."
    if "403" in msg or "PERMISSION_DENIED" in msg or "API key" in msg or "400" in msg:
        return "⚠️ Hay un problema con la clave de la IA. Avisa al administrador."
    return "⚠️ Ocurrió un error inesperado. Intenta de nuevo en unos momentos."


# ============================================================
# CARGA SILENCIOSA DE DOCUMENTOS
# ============================================================
doc_files = list_document_files()
docs_context = load_documents(files_signature(doc_files)) if doc_files else ""

# ============================================================
# ESTADO
# ============================================================
if "active_section" not in st.session_state:
    st.session_state.active_section = "assistant"

if "messages" not in st.session_state:
    st.session_state.messages = [
        {
            "role": "assistant",
            "content": "Hola. Soy el **asistente de Elite Flower**. 🌸\n\n¿En qué te puedo ayudar?",
        }
    ]

# ============================================================
# CABECERA
# ============================================================
st.markdown(
    '<div class="topbar">'
    '<div class="topbar-brand">ELITE FLOWER <span>●</span> WORKSPACE</div>'
    '<div class="topbar-note">Asistente y operaciones geográficas</div>'
    '</div>',
    unsafe_allow_html=True,
)

left, assistant_nav, gateways_nav, right = st.columns([3, 1.35, 1.35, 3], gap="small")
with assistant_nav:
    if st.button(
        "✦ Asistente",
        key="nav_assistant",
        type="primary" if st.session_state.active_section == "assistant" else "secondary",
        width="stretch",
    ):
        st.session_state.active_section = "assistant"
        st.rerun()
with gateways_nav:
    if st.button(
        "◉ Gateways",
        key="nav_gateways",
        type="primary" if st.session_state.active_section == "gateways" else "secondary",
        width="stretch",
    ):
        st.session_state.active_section = "gateways"
        st.rerun()

if st.session_state.active_section == "gateways":
    render_gateways()
    st.stop()

st.markdown(
    "<style>.block-container { max-width: 820px; }</style>"
    '<div class="hero">'
    '<div class="hero-title">ELITE FLOWER <span class="neon">AI ASSISTANT</span></div>'
    '<div class="hero-sub">Pregúntame lo que necesites</div>'
    "</div>",
    unsafe_allow_html=True,
)

# ============================================================
# CHAT
# ============================================================
for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

prompt = st.chat_input("Escribe tu pregunta...")

if prompt:
    st.session_state.messages.append({"role": "user", "content": prompt})
    with st.chat_message("user"):
        st.markdown(prompt)

    with st.chat_message("assistant"):
        if not API_KEY:
            answer = "⚠️ El asistente no está configurado todavía. Avisa al administrador."
            st.markdown(answer)
        else:
            with st.spinner("Pensando..."):
                try:
                    answer = ask_gemini(st.session_state.messages, docs_context)
                except Exception as e:
                    answer = friendly_error(e)
            st.markdown(answer)

    st.session_state.messages.append({"role": "assistant", "content": answer})
