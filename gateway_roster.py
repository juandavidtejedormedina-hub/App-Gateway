"""Cruza el listado de fincas del Excel con los puntos del KMZ, sin inventar ubicaciones."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
import unicodedata
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from gateway_data import REGIONS


MAX_XLSX_BYTES = 20 * 1024 * 1024
ZONE_COLORS = {
    "Antioquia": "#a78bfa",
    "Cachipay": "#ffad7c",
    "Centro": "#59e6d5",
    "Norte": "#77b9ff",
    "Occidente": "#ffd166",
    "Sur": "#ff83bc",
    "Medicion de la linea": "#8da3b8",
}


def normalized_name(value: object) -> str:
    """Normaliza grafías inequívocas; no equipara variedades o sedes distintas."""
    text = unicodedata.normalize("NFKD", str(value).casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    words = re.findall(r"[a-z0-9]+", text)
    while words and words[0] in {"finca", "el", "la", "las"}:
        words.pop(0)
    return "".join(words)


def load_roster(path: str | Path) -> list[dict]:
    """Lee las fincas en el orden de la hoja Detallado, sin alterar el libro."""
    workbook_path = Path(path).expanduser().resolve(strict=True)
    if workbook_path.suffix.lower() != ".xlsx" or workbook_path.stat().st_size > MAX_XLSX_BYTES:
        raise ValueError("El listado debe ser un XLSX de hasta 20 MB")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, ParseError) as error:
        raise ValueError("No se pudo leer el Excel") from error
    try:
        if "Detallado" not in workbook.sheetnames:
            raise ValueError("El Excel no contiene la hoja Detallado")
        rows: list[dict] = []
        current_zone: str | None = None
        found_header = False
        zone_positions = {zone: 0 for zone in REGIONS}
        for values in workbook["Detallado"].iter_rows(values_only=True):
            first = str(values[0] or "").strip() if values else ""
            second = str(values[1] or "").strip() if len(values) > 1 else ""
            if first.casefold() == "zona" and second.casefold() == "fincas":
                found_header = True
                continue
            if not found_header:
                continue
            if first in REGIONS:
                current_zone = first
            elif first:
                current_zone = None  # Totales y encabezados no son fincas.
            if current_zone and second:
                zone_positions[current_zone] += 1
                rows.append({"zone": current_zone, "name": second[:120],
                             "order": len(rows) + 1,
                             "zone_order": zone_positions[current_zone]})
            if len(rows) > 5000:
                raise ValueError("La hoja Detallado supera 5000 fincas")
        if not rows:
            raise ValueError("La hoja Detallado no contiene fincas por zona")
        return rows
    finally:
        workbook.close()


def reconcile_features(features: list[dict], roster: list[dict]) -> tuple[list[dict], list[dict]]:
    """Solo vincula nombres equivalentes dentro de la misma zona.

    Los nombres genéricos del KMZ (p. ej. una finca con varias variedades en el
    Excel) siguen visibles como referencias, pero no se asignan por suposición.
    """
    output = deepcopy(features)
    entries = [{**row, "on_map": False, "map_name": None} for row in roster]
    roster_index = {(row["zone"], normalized_name(row["name"])): row
                    for row in entries}

    for feature in output:
        props = feature["properties"]
        zone = str(props.get("region") or "")
        props["zone_color"] = ZONE_COLORS.get(zone, "#a8b8c5")
        if props.get("kind") != "finca":
            continue
        key = (zone, normalized_name(props.get("name") or ""))
        match = roster_index.get(key)
        if match is None:
            props["kind"] = "referencia_punto"
            props["in_roster"] = False
            continue
        match["on_map"] = True
        match["map_name"] = props["name"]
        props["in_roster"] = True
        props["roster_name"] = match["name"]
    return output, entries
